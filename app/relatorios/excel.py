"""Exportação da análise para Excel."""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import BinaryIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.engine.cenarios import Analise
from app.relatorios import agregacao

CABECALHO = PatternFill("solid", fgColor="1F3864")
FONTE_CAB = Font(color="FFFFFF", bold=True)


def _aba(wb: Workbook, titulo: str, linhas: list[dict]):
    ws = wb.create_sheet(titulo)
    if not linhas:
        ws.append(["Sem dados"])
        return
    colunas = list(linhas[0].keys())
    ws.append([c.replace("_", " ") for c in colunas])
    for cel in ws[1]:
        cel.fill, cel.font = CABECALHO, FONTE_CAB
        cel.alignment = Alignment(wrap_text=True, vertical="center")
    for l in linhas:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in (l.get(c) for c in colunas)])
    for idx, c in enumerate(colunas, start=1):
        letra = get_column_letter(idx)
        largura = max(len(c), *(len(str(l.get(c) or "")) for l in linhas[:200]))
        ws.column_dimensions[letra].width = min(max(10, largura + 2), 60)
        fmt = "0.00" if c.endswith("_pct") else "#,##0.00"
        if any(isinstance(l.get(c), Decimal) for l in linhas[:50]):
            for cel in ws[letra][1:]:
                cel.number_format = fmt
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def exportar(a: Analise, destino: Path | BinaryIO, extras: dict[str, list[dict]] | None = None) -> Path | BinaryIO:
    wb = Workbook()
    wb.remove(wb.active)
    _aba(wb, "Resumo", agregacao.resumo(a))
    _aba(wb, "Por produto (saídas)", agregacao.por_produto(a))
    _aba(wb, "Por fornecedor (entradas)", agregacao.por_fornecedor(a))
    _aba(wb, "Classificação", agregacao.classificacao(a))
    _aba(wb, "Itens", agregacao.itens(a))
    _aba(wb, "Premissas", agregacao.premissas(a))
    for titulo, linhas in (extras or {}).items():
        _aba(wb, titulo, linhas)
    _aba(wb, "Alertas", [{"alerta": t} for t in a.alertas] or [{"alerta": "Nenhum"}])
    if isinstance(destino, Path):
        destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def exportar_negociacao(compras, ano: int, destino: Path | BinaryIO) -> Path | BinaryIO:
    """Planilha para negociação: resumo por fornecedor e compras por fornecedor x produto, com preço de nota de
    equilíbrio e o melhor fornecedor do mesmo produto a preços de hoje."""
    from app.relatorios import negociacao

    wb = Workbook()
    wb.remove(wb.active)
    _aba(wb, "Fornecedores", [{
        "cnpj": f.cnpj, "fornecedor": f.nome, "regime": f.regime, "uf": f.uf, "produtos": f.produtos,
        "notas": f.notas, "compras_preco_de_nota": f.preco_nota_atual, "custo_efetivo_atual": f.custo_atual,
        f"custo_{ano}_a_precos_de_hoje": f.custo_precos_hoje, "variacao_custo": f.variacao_custo,
        "variacao_custo_pct": f.variacao_custo_pct, f"custo_{ano}_com_repasse": f.custo_ano,
        f"preco_de_nota_equilibrio_{ano}": f.preco_nota_equilibrio,
        "variacao_preco_de_nota_equilibrio_pct": f.variacao_preco_nota_equilibrio_pct,
        "economia_trocando_fornecedor": f.economia_trocando,
    } for f in negociacao.fornecedores(compras)])
    _aba(wb, "Fornecedor x produto", [{
        "cnpj": c.fornecedor, "fornecedor": c.fornecedor_nome, "regime": c.regime, "ean": c.gtin,
        "codigos": ", ".join(c.codigos), "descricao": c.descricao, "ncm": c.ncm, "cclasstrib": c.cclasstrib,
        "unidade": c.unidade, "quantidade": c.quantidade, "preco_de_nota_unit": c.preco_nota_unit,
        "custo_unit_atual": c.custo_unit_atual, f"custo_unit_{ano}_a_precos_de_hoje": c.custo_unit_precos_hoje,
        "variacao_custo_pct": c.variacao_custo_precos_hoje_pct, f"custo_unit_{ano}_com_repasse": c.custo_unit_ano,
        f"preco_de_nota_equilibrio_unit_{ano}": c.preco_nota_equilibrio_unit,
        "variacao_preco_de_nota_equilibrio_pct": c.variacao_preco_nota_equilibrio_pct,
        "equilibrio_sobre_valor_liquido_pct": c.preco_equilibrio_pct,
        f"imposto_seletivo_{ano}": c.imposto_seletivo_ano, "tratamento_is": c.tratamento_is,
        "melhor_fornecedor_precos_de_hoje": c.melhor.fornecedor_nome if c.melhor and c.melhor is not c else "",
        f"melhor_custo_unit_{ano}": c.melhor.custo_unit_precos_hoje if c.melhor and c.melhor is not c else None,
        "economia_trocando": c.economia_trocando,
    } for c in sorted(compras, key=lambda c: (c.fornecedor_nome, c.descricao))])
    if isinstance(destino, Path):
        destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino
