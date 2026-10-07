from decimal import Decimal as D
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import db, servicos
from app.api import main
from app.engine.cenarios import Empresa, analisar
from app.engine.premissas import Premissas
from app.ingest.arquivos import ler_conteudos
from app.relatorios import negociacao
from tests.conftest import CNPJ_EMPRESA, FX, CalculadoraFake

NOTEBOOK = "7890000000017"


def _notebook_do_simples() -> bytes:
    """O mesmo notebook (mesmo EAN) comprado de um fornecedor do Simples por R$ 950."""
    xml = (FX / "xml" / "entrada_fornecedor_simples.xml").read_text(encoding="utf-8")
    return (xml.replace("35261077888999000155550010000005671000056789", "35261077888999000155550010000005681000056780")
               .replace("<cProd>POTE-10</cProd>", f"<cProd>NOTE-15</cProd><cEAN>{NOTEBOOK}</cEAN>")
               .replace("POTE PLASTICO 10L", "NOTEBOOK 15 POLEGADAS").replace("<NCM>39241000</NCM>", "<NCM>84713012</NCM>")
               .replace("<qCom>50.0000</qCom><vProd>500.00</vProd>", "<qCom>1.0000</qCom><vProd>950.00</vProd>")
               .replace("<vCredICMSSN>12.50</vCredICMSSN>", "<vCredICMSSN>0.00</vCredICMSSN>")).encode()


def _arquivos():
    return [(p.name, p.read_bytes()) for p in sorted((FX / "xml").glob("*.xml"))] + [("simples-nb.xml", _notebook_do_simples())]


@pytest.fixture
def analise():
    docs = ler_conteudos(_arquivos()).documentos
    return analisar(Empresa(CNPJ_EMPRESA, "LOJA", "real", "SP", 3550308), docs, Premissas(anos=[2027, 2033]),
                    CalculadoraFake(), {"ean:7890000000024": ("200", "200003")})


def test_custos_equilibrio_e_melhor_fornecedor(analise):
    compras = {(c.fornecedor, c.chave): c for c in negociacao.compras(analise, 2033)}
    regular = compras[("44555666000177", f"ean:{NOTEBOOK}")]
    simples = compras[("77888999000155", f"ean:{NOTEBOOK}")]
    # Regime normal: hoje 1.000 de nota, custo 789,52. Em 2033, se mantiver a nota em 1.000, tudo vira valor líquido
    # (sem ICMS/IPI/PIS/COFINS) e o custo é 1.000; com repasse, 744,15. Equilíbrio: nota de 789,52 (-21,05%).
    assert regular.custo_unit_atual == D("789.5200") and regular.custo_unit_ano == D("744.1500")
    assert regular.custo_unit_precos_hoje == D("1000.0000") and regular.variacao_custo_precos_hoje_pct == D("26.66")
    assert regular.preco_nota_equilibrio_unit == D("789.5200") and regular.variacao_preco_nota_equilibrio_pct == D("-21.05")
    # Simples: preço não muda (DAS), crédito cai de 9,25% para 4% (premissa).
    assert simples.custo_unit_atual == D("862.1200") and simples.custo_unit_precos_hoje == D("912.0000")
    assert simples.variacao_preco_nota_equilibrio_pct == D("-5.47")
    # A preços de hoje, em 2033 o fornecedor do Simples fica mais barato que o do regime normal
    assert regular.melhor is simples and simples.melhor is simples
    assert regular.economia_trocando == D("88.00")
    pote = compras[("77888999000155", "cod:77888999:POTE-10")]
    assert pote.melhor is None and pote.economia_trocando is None  # sem outro fornecedor


def test_resumos(analise):
    lista = negociacao.compras(analise, 2033)
    forn = {f.cnpj: f for f in negociacao.fornecedores(lista)}
    assert forn["44555666000177"].economia_trocando == D("88.00") and forn["77888999000155"].produtos == 2
    assert list(forn)[0] == "44555666000177"                      # maior aumento de custo (a preços de hoje) primeiro
    prod = {p.chave: p for p in negociacao.produtos(analise, lista, 2033)}
    nb = prod[f"ean:{NOTEBOOK}"]
    assert nb.fornecedores == 2 and nb.vendas_liquidas == D("1197.90")
    assert nb.variacao_preco_venda_pct == D("1.02")
    assert prod["ean:7890000000024"].carga_venda_ano_pct == D("0.00")   # arroz: cesta básica


def test_telas_excel_e_cache(banco, monkeypatch):
    monkeypatch.setattr(main, "calculadora", CalculadoraFake())
    with db.sessao() as s:
        e = servicos.criar_empresa(s, CNPJ_EMPRESA, "LOJA", "real", "SP", 3550308)
        servicos.importar_arquivos(s, e, [(p.name, p.read_bytes()) for p in sorted((FX / "xml").glob("*.xml"))])
        id_ = e.id
    c = TestClient(main.app)
    lista = c.get(f"/empresas/{id_}/fornecedores?ano=2033")
    assert lista.status_code == 200 and "UTILIDADES PEQUENA ME" in lista.text and "Preço de nota de equilíbrio" in lista.text
    assert "MEI" in c.get(f"/empresas/{id_}/fornecedores?ano=2033&regime=simples").text  # filtro renderiza
    assert "sem outro fornecedor" in c.get(f"/empresas/{id_}/fornecedores/44555666000177?ano=2033").text

    # nova nota: o cache da análise precisa ser invalidado e o melhor fornecedor aparecer
    with db.sessao() as s:
        servicos.importar_arquivos(s, s.get(servicos.models.Empresa, id_), [("simples-nb.xml", _notebook_do_simples())])
    detalhe = c.get(f"/empresas/{id_}/fornecedores/44555666000177?ano=2033")
    assert "UTILIDADES PEQUENA ME" in detalhe.text and "88,00" in detalhe.text
    produto = c.get(f"/empresas/{id_}/produtos/ean:{NOTEBOOK}?ano=2033")
    assert produto.status_code == 200 and "melhor em 2033" in produto.text
    assert "NOTEBOOK" in c.get(f"/empresas/{id_}/produtos?ano=2027").text
    assert c.get(f"/empresas/{id_}/fornecedores/00000000000000").status_code == 404

    xlsx = c.get(f"/empresas/{id_}/negociacao.xlsx?ano=2033")
    wb = load_workbook(BytesIO(xlsx.content))
    assert wb.sheetnames == ["Fornecedores", "Fornecedor x produto"]
    cab = [c.value for c in wb["Fornecedor x produto"][1]]
    assert "melhor fornecedor precos de hoje" in cab and "preco de nota equilibrio unit 2033" in cab
