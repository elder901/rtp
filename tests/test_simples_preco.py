from decimal import Decimal as D

from fastapi.testclient import TestClient

from app import db, servicos
from app.api import main
from app.engine.cenarios import Empresa, analisar
from app.engine.premissas import Premissas
from app.engine.simples import comparar
from app.ingest.arquivos import ler_caminho
from app.relatorios import agregacao
from tests.conftest import CNPJ_EMPRESA, FX, CalculadoraFake


def _analise(regime: str):
    docs = ler_caminho(FX / "xml").documentos
    return analisar(Empresa(CNPJ_EMPRESA, "LOJA", regime, "SP", 3550308), docs, Premissas(anos=[2027, 2033]),
                    CalculadoraFake(), {"V-ARZ5": ("200", "200003")})


def test_simples_x_regular():
    linhas = {l["ano"]: l for l in comparar(_analise("simples"), _analise("presumido"))}
    l = linhas[2033]
    assert l["simples_carga"] == D("132.00")                      # 6% de DAS sobre 2.200,00 vendidos
    assert l["hibrido_das"] == D("66.66")                         # DAS sem PIS/COFINS (15,5%) e ICMS (34%)
    assert l["hibrido_debitos"] == D("317.45")                    # CBS/IBS do notebook; arroz com redução de 100%
    assert l["hibrido_creditos"] == D("197.20") + D("20.00")      # fornecedor regular + 4% do fornecedor do Simples
    assert l["hibrido_carga"] == D("66.66") + D("317.45") - D("217.20")
    assert l["hibrido_credito_clientes_pj"] == D("317.45")
    assert l["melhor"] == "Simples puro"
    # 2027: ICMS ainda existe, só a parcela de PIS/COFINS sai do DAS
    assert linhas[2027]["hibrido_das"] == (D("132.00") * D("0.845")).quantize(D("0.01"))


def test_preco_neutro_com_margem():
    prod = {p["codigo"]: p for p in agregacao.por_produto(_analise("real"))}
    nb = prod["V-NB01"]
    assert nb["gtin"] == "7890000000017"
    assert nb["custo_unit_atual"] == D("789.5200")                # 1.050,00 pagos - 260,48 de créditos
    assert nb["margem_unit"] == D("408.38")                       # 1.197,90 líquido - 789,52
    assert nb["var_preco_2033_pct"] == D("1.02")                  # receita líquida constante
    assert nb["preco_neutro_2033_pct"] == D("-2.80")              # margem constante: custo cai com crédito amplo
    assert prod["V-ARZ5"]["preco_neutro_2033_pct"] is not None
    assert "BRINDE" not in prod


def test_painel_empresa_simples(banco, monkeypatch):
    monkeypatch.setattr(main, "calculadora", CalculadoraFake())
    with db.sessao() as s:
        e = servicos.criar_empresa(s, CNPJ_EMPRESA, "LOJA", "simples", "SP", 3550308)
        servicos.importar_arquivos(s, e, [(p.name, p.read_bytes()) for p in (FX / "xml").glob("*.xml")])
        id_ = e.id
    c = TestClient(main.app)
    painel = c.get(f"/empresas/{id_}/analise")
    assert "permanecer no Simples" in painel.text and "Preço neutro" in painel.text
    excel = c.get(f"/empresas/{id_}/analise/excel")
    assert excel.status_code == 200
    from io import BytesIO

    from openpyxl import load_workbook
    assert "Simples x regular" in load_workbook(BytesIO(excel.content)).sheetnames
